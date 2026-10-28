"""
Lecture des fichiers Valorix (Excel ou PDF) et des documents « Autres »
(Excel, CSV ou PDF), puis enregistrement dans un lot d'import.

Toutes les fonctions de lecture retournent des lignes (dict) prêtes à être
affichées dans un tableau de vérification ; rien n'est écrit en base.
"""

import csv
import io
import uuid
from collections import Counter
from datetime import datetime
from decimal import Decimal

from django.db import transaction as db_transaction
from django.utils import timezone

from importation.models import LotImport, SourceImport
from importation.outils import (
    a_zero,
    convertir_date,
    convertir_montant,
    marquer_doublons,
    montant_positif,
    montant_vers_texte,
    normaliser_entete,
    normaliser_texte,
)
from importation.pdf import (
    ExtractionError,
    extraire_par_blocs,
    lire_pages_pdf,
    regrouper_pages,
)
from valorix.models import TransactionValorix, TypeTransactionCaisse
from valorix.prompts import PROMPT_DOCUMENT_AUTRE, PROMPT_VALORIX
from valorix.schemas import ExtractionDocumentAutre, ExtractionValorix

T = TypeTransactionCaisse

# Libellés de catégories (normalisés) -> type
CATEGORIES = {
    "depot cheques": T.DEPOT_CHEQUE,
    "retrait cheques": T.RETRAIT_CHEQUE,
    "others": T.AUTRE,
    "depot especes": T.DEPOT_ESPECES,
    "retrait especes": T.RETRAIT_ESPECES,
    "retrait especes carte visa cgf": T.RETRAIT_ESPECES,
    "transferts": T.TRANSFERT,
    "depot virements": T.DEPOT_VIREMENTS,
    "retrait virements": T.RETRAIT_VIREMENTS,
}

# Colonnes du fichier Excel Valorix (indices à partir de 0)
COL_TRANSACTION, COL_DATE, COL_COMPTE, COL_NOM = 1, 4, 8, 9
COL_CREDIT, COL_DEBIT, COL_REFERENCES = 12, 13, 14
PREMIERE_LIGNE_EXCEL = 13


# ------------------------------------------------------------------
# Construction d'une ligne commune
# ------------------------------------------------------------------

def _verifier_montants(debit, credit, erreurs, alertes):
    if (debit is not None and debit < 0) or (credit is not None and credit < 0):
        erreurs.append("Les montants débit/crédit ne doivent pas être négatifs.")
    debit, credit = montant_positif(debit), montant_positif(credit)
    if debit and credit:
        erreurs.append("Débit et crédit sont renseignés simultanément.")
    elif not debit and not credit:
        alertes.append("Montant nul : cette ligne ne pourra pas être rapprochée.")
    return debit, credit


def construire_ligne(
    index, numero_transaction, date_transaction, numero_compte, nom_complet,
    debit, credit, references, types, erreurs=None,
):
    erreurs = list(erreurs or [])
    alertes = []
    debit, credit = _verifier_montants(debit, credit, erreurs, alertes)

    return {
        "index": index,
        "numero_transaction": numero_transaction,
        "date_transaction": (
            date_transaction.isoformat() if date_transaction else None
        ),
        "numero_compte": numero_compte,
        "nomComplet": nom_complet,
        "debit": montant_vers_texte(debit),
        "credit": montant_vers_texte(credit),
        "references": references,
        "types": types,
        "doublon_probable": False,
        "erreurs": erreurs,
        "alertes": alertes,
    }


def _texte(valeur):
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return str(valeur if valeur is not None else "").strip()


# ------------------------------------------------------------------
# Relevé Valorix — Excel
# ------------------------------------------------------------------

def lire_excel_valorix(fichier):
    from openpyxl import load_workbook

    try:
        fichier.seek(0)
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError("Le fichier Excel est invalide ou illisible.") from exc

    lignes = []
    try:
        type_courant = None
        derniere = max(
            COL_TRANSACTION, COL_DATE, COL_COMPTE, COL_NOM,
            COL_CREDIT, COL_DEBIT, COL_REFERENCES,
        )

        for numero_ligne, ligne in enumerate(
            classeur.active.iter_rows(min_row=PREMIERE_LIGNE_EXCEL, values_only=True),
            start=PREMIERE_LIGNE_EXCEL,
        ):
            if not any(v is not None for v in ligne):
                continue

            if len(ligne) <= derniere:
                lignes.append(
                    construire_ligne(
                        len(lignes) + 1, "", None, "", "", None, None, "",
                        type_courant or T.AUTRE,
                        [f"Ligne Excel {numero_ligne} : colonnes manquantes."],
                    )
                )
                continue

            numero = ligne[COL_TRANSACTION]

            if isinstance(numero, str):  # en-tête de catégorie ou libellé
                categorie = normaliser_texte(numero)
                categorie = " ".join(categorie.replace("/", " ").split())
                if categorie in CATEGORIES:
                    type_courant = CATEGORIES[categorie]
                continue

            if not isinstance(numero, (int, float)):
                continue  # ligne de synthèse

            erreurs = []
            if isinstance(numero, float) and not numero.is_integer():
                erreurs.append("Le numéro de transaction doit être entier.")
            if type_courant is None:
                erreurs.append("Type de transaction non reconnu.")

            date_tx = debit = credit = None
            try:
                date_tx = convertir_date(ligne[COL_DATE], classeur.epoch)
            except ValueError as exc:
                erreurs.append(str(exc))
            try:
                debit = convertir_montant(ligne[COL_DEBIT])
                credit = convertir_montant(ligne[COL_CREDIT])
            except ValueError as exc:
                erreurs.append(str(exc))

            lignes.append(
                construire_ligne(
                    len(lignes) + 1,
                    str(int(numero)),
                    date_tx,
                    _texte(ligne[COL_COMPTE]),
                    _texte(ligne[COL_NOM]),
                    debit, credit,
                    _texte(ligne[COL_REFERENCES]),
                    type_courant or T.AUTRE,
                    [f"Ligne Excel {numero_ligne} : {e}" for e in erreurs],
                )
            )
    finally:
        classeur.close()

    return lignes


# ------------------------------------------------------------------
# Relevé Valorix — PDF
# ------------------------------------------------------------------

def lire_pdf_valorix(fichier):
    pages = lire_pages_pdf(fichier)
    blocs = regrouper_pages(pages, pages_par_bloc=3)
    brutes = extraire_par_blocs(
        blocs, ExtractionValorix, PROMPT_VALORIX, "relevé Valorix", max_workers=3
    )

    lignes = []
    for i, brut in enumerate(brutes, 1):
        erreurs = []
        libelle = normaliser_texte(brut.type_transaction).replace("/", " ")
        libelle = " ".join(libelle.split())
        type_tx = CATEGORIES.get(libelle)
        if type_tx is None:
            erreurs.append(f"Type de transaction non reconnu : {brut.type_transaction!r}.")

        date_tx = brut.date_transaction
        if date_tx is None:
            erreurs.append("Date de transaction manquante.")
        elif timezone.is_naive(date_tx):
            date_tx = timezone.make_aware(date_tx, timezone.get_current_timezone())

        lignes.append(
            construire_ligne(
                i, _texte(brut.numero_transaction), date_tx,
                _texte(brut.numero_compte), _texte(brut.nomComplet),
                brut.debit, brut.credit, _texte(brut.references),
                type_tx or T.AUTRE, erreurs,
            )
        )
    return lignes


def lire_releve_valorix(fichier):
    nom = fichier.name.lower()
    if nom.endswith(".xlsx"):
        return lire_excel_valorix(fichier)
    if nom.endswith(".pdf"):
        return lire_pdf_valorix(fichier)
    raise ExtractionError("Format non pris en charge : utilisez un fichier .xlsx ou .pdf.")


# ------------------------------------------------------------------
# Document « Autres » : numéro compte client / date / montant net
# ------------------------------------------------------------------

COLONNES_REQUISES = {
    "compte": "Numéro compte client",
    "date": "Date transaction",
    "montant": "Montant net",
}

_EXACT = {
    "compte": {
        "numero compte client", "numero de compte client", "n compte client",
        "no compte client", "compte client", "numero compte",
        "numero de compte", "n compte", "num compte", "compte",
    },
    "date": {"date transaction", "date de transaction", "date operation", "date"},
    "montant": {"montant net", "net"},
}


def detecter_colonnes(entetes):
    """
    Associe les en-têtes d'un fichier aux 3 colonnes requises.
    Retourne (colonnes, manquantes) où colonnes = {"compte": idx, ...}.
    """
    normalises = [normaliser_entete(e) for e in entetes]
    colonnes, utilises = {}, set()

    def attribuer(cle, predicats):
        for predicat in predicats:
            for i, texte in enumerate(normalises):
                if i not in utilises and texte and predicat(texte):
                    colonnes[cle] = i
                    utilises.add(i)
                    return

    attribuer("montant", [
        lambda t: t in _EXACT["montant"],
        lambda t: "montant" in t and "net" in t,
        lambda t: "montant" in t,
    ])
    attribuer("compte", [
        lambda t: t in _EXACT["compte"],
        lambda t: "compte" in t,
    ])
    attribuer("date", [
        lambda t: t in _EXACT["date"],
        lambda t: t.startswith("date"),
        lambda t: "date" in t,
    ])

    manquantes = [
        libelle for cle, libelle in COLONNES_REQUISES.items() if cle not in colonnes
    ]
    return colonnes, manquantes


def _erreur_colonnes(manquantes, entetes):
    trouvees = ", ".join(str(e) for e in entetes if e) or "aucune"
    return ExtractionError(
        f"Colonnes manquantes : {', '.join(manquantes)}. "
        f"Colonnes détectées dans le fichier : {trouvees}."
    )


def _lignes_depuis_tableau(lignes_brutes, sens, epoch=None):
    """lignes_brutes : liste de listes (la 1re ligne contenant les 3 colonnes = en-têtes)."""
    entetes, debut, colonnes = None, None, None
    for i, ligne in enumerate(lignes_brutes[:30]):
        cols, manquantes = detecter_colonnes(ligne)
        if not manquantes:
            entetes, debut, colonnes = ligne, i + 1, cols
            break

    if colonnes is None:
        premiere = next((l for l in lignes_brutes[:30] if any(l)), [])
        cols, manquantes = detecter_colonnes(premiere)
        raise _erreur_colonnes(manquantes, premiere)

    lignes = []
    for ligne in lignes_brutes[debut:]:
        def cellule(cle):
            i = colonnes[cle]
            return ligne[i] if i < len(ligne) else None

        compte, date_val, montant = cellule("compte"), cellule("date"), cellule("montant")
        if all(v in (None, "") for v in (compte, date_val, montant)):
            continue
        if date_val in (None, "") and (
            _texte(compte) == ""
            or normaliser_entete(compte).startswith(("total", "sous total"))
        ):
            continue  # ligne de total / commentaire

        lignes.append(
            construire_ligne_autre(len(lignes) + 1, compte, date_val, montant, sens, epoch)
        )
    return lignes


def construire_ligne_autre(index, compte, date_val, montant, sens, epoch=None):
    erreurs = []
    compte_txt = _texte(compte)
    if not compte_txt:
        erreurs.append("Numéro de compte client manquant.")

    date_tx = None
    try:
        date_tx = convertir_date(date_val, epoch)
    except ValueError as exc:
        erreurs.append(str(exc))

    valeur = None
    try:
        valeur = convertir_montant(montant)
        if valeur is None:
            erreurs.append("Montant net manquant.")
        elif valeur < 0:
            erreurs.append("Montant net négatif non accepté.")
    except ValueError as exc:
        erreurs.append(str(exc))

    debit = credit = None
    if valeur is not None and valeur >= 0:
        if sens == "DEBIT":
            debit = valeur
        else:
            credit = valeur

    return construire_ligne(
        index, f"DOC-{index:05d}", date_tx, compte_txt, "", debit, credit, "",
        T.DOCUMENT_AUTRE, erreurs,
    )


def lire_document_autre(fichier, sens):
    if sens not in ("DEBIT", "CREDIT"):
        raise ExtractionError("Le sens des montants doit être DEBIT ou CREDIT.")

    nom = fichier.name.lower()
    fichier.seek(0)

    if nom.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            classeur = load_workbook(fichier, read_only=True, data_only=True)
        except Exception as exc:
            raise ExtractionError("Le fichier Excel est invalide ou illisible.") from exc
        try:
            brutes = [list(l) for l in classeur.active.iter_rows(values_only=True)]
            return _lignes_depuis_tableau(brutes, sens, classeur.epoch)
        finally:
            classeur.close()

    if nom.endswith(".csv"):
        contenu = fichier.read()
        try:
            texte = contenu.decode("utf-8-sig")
        except UnicodeDecodeError:
            texte = contenu.decode("latin-1")
        try:
            dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=",;\t|")
        except csv.Error:
            dialecte = csv.excel
            dialecte.delimiter = ";"
        brutes = [list(l) for l in csv.reader(io.StringIO(texte), dialecte)]
        return _lignes_depuis_tableau(brutes, sens)

    if nom.endswith(".pdf"):
        pages = lire_pages_pdf(fichier)
        blocs = regrouper_pages(pages, pages_par_bloc=3)
        brutes = extraire_par_blocs(
            blocs, ExtractionDocumentAutre, PROMPT_DOCUMENT_AUTRE,
            "document de transactions", max_workers=3,
        )
        return [
            construire_ligne_autre(
                i, b.numero_compte_client, b.date_transaction, b.montant_net, sens
            )
            for i, b in enumerate(brutes, 1)
        ]

    raise ExtractionError("Format non pris en charge : utilisez .xlsx, .csv ou .pdf.")


def generer_modele_document_autre():
    """Modèle Excel à télécharger : en-têtes + une ligne d'exemple."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(list(COLONNES_REQUISES.values()))
    ws.append(["0012345678", datetime.now().strftime("%d/%m/%Y"), 150000])
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="1F3A5F")
    for lettre, largeur in zip("ABC", (26, 20, 18)):
        ws.column_dimensions[lettre].width = largeur
    ws["E1"] = "Supprimez la ligne d'exemple avant l'import."
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


# ------------------------------------------------------------------
# Doublons
# ------------------------------------------------------------------

def _date_locale(valeur):
    if isinstance(valeur, str):
        valeur = datetime.fromisoformat(valeur)
    if timezone.is_aware(valeur):
        valeur = timezone.localtime(valeur)
    return valeur.date()


def cle_transaction(types, numero_transaction, date_tx, numero_compte, debit, credit):
    base = (_date_locale(date_tx), a_zero(debit), a_zero(credit), (numero_compte or "").strip())
    if types == T.DOCUMENT_AUTRE:
        return ("DOC",) + base
    return ("VLX", (numero_transaction or "").strip()) + base


def _compteur_existant(dates, document_autre):
    if not dates:
        return Counter()
    qs = TransactionValorix.objects.filter(
        date_transaction__date__range=(min(dates), max(dates))
    )
    qs = qs.filter(types=T.DOCUMENT_AUTRE) if document_autre else qs.exclude(types=T.DOCUMENT_AUTRE)
    return Counter(
        cle_transaction(types, num, date_tx, compte, debit, credit)
        for types, num, date_tx, compte, debit, credit in qs.values_list(
            "types", "numero_transaction", "date_transaction",
            "numero_compte", "debit", "credit",
        )
    )


def marquer_doublons_lignes(lignes, document_autre=False):
    """Renseigne ``doublon_probable`` sur des lignes de prévisualisation."""
    valides = [l for l in lignes if l["date_transaction"]]
    compteur = _compteur_existant(
        [_date_locale(l["date_transaction"]) for l in valides], document_autre
    )
    cles = [
        cle_transaction(
            l["types"], l["numero_transaction"], l["date_transaction"],
            l["numero_compte"],
            Decimal(l["debit"]) if l["debit"] else None,
            Decimal(l["credit"]) if l["credit"] else None,
        )
        for l in valides
    ]
    for ligne, doublon in zip(valides, marquer_doublons(cles, compteur)):
        ligne["doublon_probable"] = doublon
    return lignes


# ------------------------------------------------------------------
# Enregistrement
# ------------------------------------------------------------------

@db_transaction.atomic
def enregistrer_import(lignes, source, nom_fichier, ignorer_doublons):
    """Enregistre des lignes validées. Retourne (lot, créées, ignorées)."""
    document_autre = source == SourceImport.DOCUMENT_AUTRE

    doublons = [False] * len(lignes)
    if ignorer_doublons:
        compteur = _compteur_existant(
            [_date_locale(l["date_transaction"]) for l in lignes], document_autre
        )
        doublons = marquer_doublons(
            [
                cle_transaction(
                    l["types"], l.get("numero_transaction"), l["date_transaction"],
                    l.get("numero_compte"), l.get("debit"), l.get("credit"),
                )
                for l in lignes
            ],
            compteur,
        )

    lot = LotImport.objects.create(source=source, nom_fichier=nom_fichier or "")
    prefixe = lot.id.hex[:6].upper()

    a_creer, ignorees = [], 0
    for ligne, doublon in zip(lignes, doublons):
        if doublon:
            ignorees += 1
            continue
        numero = ligne.get("numero_transaction") or ""
        if document_autre:
            numero = f"DOC-{prefixe}-{len(a_creer) + 1:05d}"
        a_creer.append(
            TransactionValorix(
                id=uuid.uuid4(),
                lot=lot,
                numero_transaction=numero,
                date_transaction=ligne["date_transaction"],
                numero_compte=ligne.get("numero_compte") or "",
                nomComplet=ligne.get("nomComplet") or "",
                debit=montant_positif(ligne.get("debit")),
                credit=montant_positif(ligne.get("credit")),
                references=ligne.get("references") or "",
                types=ligne["types"],
            )
        )

    TransactionValorix.objects.bulk_create(a_creer, batch_size=500)

    lot.nombre_lignes = len(a_creer)
    lot.nombre_doublons_ignores = ignorees
    lot.save(update_fields=["nombre_lignes", "nombre_doublons_ignores"])
    return lot, len(a_creer), ignorees
